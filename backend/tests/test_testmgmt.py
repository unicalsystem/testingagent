"""CSV/Excel import, column mapping, test runs and result recording."""
from __future__ import annotations

import csv
import io
import json

import openpyxl

from testpilot import importer as im
from testpilot import worker

from .conftest import allow

HEAD = ["Test Case ID (Doc)", "Test Case Title", "Module", "Sub-module", "Priority", "Procedure / Steps",
        "Expected Result", "Result", "Executed By", "Executed On", "Assigned To", "Risk Area"]
ROWS = [
    ["CON-LOG-001", "Login with valid credentials", "Login & Access", "Login", "high", "1. Open.\n2. Sign in.",
     "Dashboard shown", "Pass", "Ravi", "2026-10-20", "Shiva", "Auth"],
    ["CON-LOG-002", "Login with invalid password", "Login & Access", "Login", "High", "1. Open.", "Error shown",
     "Not Run", "", "", "Shiva", "Auth"],
    ["CON-DSH-001", "Dashboard widgets load", "Dashboard", "Widgets", "Medium", "1. Open dashboard.",
     "Widgets visible", "Fail", "Ravi", "20/10/2026", "Anita", ""],
    ["", "", "", "", "", "", "", "", "", "", "", ""],                       # blank row: ignored
    ["", "", "Dashboard", "", "Low", "", "", "", "", "", "", ""],           # no title or id: error
    ["CON-LOG-001", "Duplicate id in file", "Login & Access", "", "", "", "", "", "", "", "", ""],
]


def csv_bytes(head=HEAD, rows=ROWS, delim=",") -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delim)
    w.writerow(head)
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode()


def mk_project(c) -> dict:
    return c.post("/api/projects", json={"name": "Connect", "base_url": "https://connect.example.com"}).json()


def preview(c, pid, data, name="cases.csv", mapping=None):
    files = {"file": (name, data, "text/csv")}
    form = {"mapping": json.dumps(mapping)} if mapping else {}
    return c.post(f"/api/projects/{pid}/imports/preview", files=files, data=form)


def commit(c, pid, data, options, name="cases.csv"):
    return c.post(f"/api/projects/{pid}/imports/commit", files={"file": (name, data, "text/csv")},
                  data={"options": json.dumps(options)})


# ---- parsing (unit) -------------------------------------------------------------------
def test_mapping_suggestions_and_normalisation():
    t = im.parse_csv(csv_bytes(delim=";"))
    m = im.suggest_mapping(t.headers)
    assert m["Test Case ID (Doc)"] == "external_id" and m["Procedure / Steps"] == "procedure"
    assert m["Risk Area"] == "attr:Risk Area"  # unknown columns are kept as custom fields
    res = im.apply_mapping(t, m)
    assert len(res.records) == 3 and res.in_file_duplicates == 1
    assert {e["row"] for e in res.errors} == {6, 7}
    r0 = res.records[0]
    assert r0.case["priority"] == "High" and r0.result["status"] == "passed" and r0.attributes["Risk Area"] == "Auth"
    assert r0.case["procedure"] == "1. Open.\n2. Sign in."


def test_template_round_trip():
    t = im.parse_xlsx(im.template_xlsx())
    assert t.sheet == "Test Cases" and t.headers[:2] == ["Test Case ID", "Test Case Title"]
    m = im.suggest_mapping(t.headers)
    assert all(v and not v.startswith("attr:") for v in m.values())  # every template column is recognised
    assert len(im.apply_mapping(t, m).records) == 2
    assert im.parse_csv(im.template_csv()).headers == t.headers


def test_old_xls_rejected_and_formula_export_is_neutralised():
    try:
        im.parse_file("a.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100)
        raise AssertionError("expected rejection")
    except im.ImportError_ as e:
        assert ".xlsx" in str(e)
    assert im._safe_cell("=HYPERLINK(1)") == "'=HYPERLINK(1)"


# ---- API ---------------------------------------------------------------------------------
def test_preview_writes_nothing_and_reports_groups(tester):
    p = mk_project(tester)
    r = preview(tester, p["id"], csv_bytes())
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["valid"] == 3 and out["error_count"] == 2 and out["new_cases"] == 3
    assert out["recommended_group_by"] == ["module"]
    assert {g["value"]: g["count"] for g in out["groups"]["module"]} == {"Login & Access": 2, "Dashboard": 1}
    assert out["results"] == {"passed": 1, "untested": 1, "failed": 1}
    assert tester.get(f"/api/projects/{p['id']}/cases").json() == []


def test_bad_mapping_is_reported(tester):
    p = mk_project(tester)
    out = preview(tester, p["id"], csv_bytes(), mapping={h: "" for h in HEAD}).json()
    assert "Title" in out["mapping_error"]
    dup = {h: "" for h in HEAD} | {"Test Case Title": "name", "Module": "name"}
    assert "both mapped" in preview(tester, p["id"], csv_bytes(), mapping=dup).json()["mapping_error"]


def test_commit_requires_consent_and_groups_runs(tester):
    p = mk_project(tester)
    mapping = im.suggest_mapping(HEAD)
    opts = {"mapping": mapping, "run_plan": {"mode": "group", "group_by": ["module"], "name": "Sprint 12"}}
    assert commit(tester, p["id"], csv_bytes(), opts).status_code == 422  # no confirm
    r = commit(tester, p["id"], csv_bytes(), {**opts, "confirm": True})
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["created"] == 3 and out["error_count"] == 2
    assert [t["name"] for t in out["test_runs"]] == ["Sprint 12 – Login & Access", "Sprint 12 – Dashboard"]
    cases = tester.get(f"/api/projects/{p['id']}/cases").json()
    assert cases[0]["external_id"] == "CON-LOG-001" and cases[0]["source"] == "import" and cases[0]["steps"] == []
    runs = {t["name"]: t for t in tester.get(f"/api/projects/{p['id']}/test-runs").json()}
    login = runs["Sprint 12 – Login & Access"]
    assert login["total"] == 2 and login["counts"]["passed"] == 1 and login["progress"] == 50
    detail = tester.get(f"/api/test-runs/{login['id']}").json()
    first = detail["items"][0]
    assert first["executed_by"] == "Ravi" and first["executed_at"].startswith("2026-10-20")
    assert first["case"]["procedure"].startswith("1. Open.")


def test_reimport_updates_instead_of_duplicating(tester):
    p = mk_project(tester)
    mapping = im.suggest_mapping(HEAD)
    commit(tester, p["id"], csv_bytes(), {"mapping": mapping, "run_plan": {"mode": "none"}, "confirm": True})
    rows = [list(ROWS[0]), list(ROWS[2])]
    rows[0][1] = "Login with valid credentials (v2)"
    assert preview(tester, p["id"], csv_bytes(rows=rows)).json()["existing_matches"] == 2
    out = commit(tester, p["id"], csv_bytes(rows=rows),
                 {"mapping": mapping, "duplicates": "update", "run_plan": {"mode": "single", "name": "Regression"},
                  "confirm": True}).json()
    assert out["created"] == 0 and out["updated"] == 2 and len(out["test_runs"]) == 1
    names = [c["name"] for c in tester.get(f"/api/projects/{p['id']}/cases").json()]
    assert len(names) == 3 and "Login with valid credentials (v2)" in names


def test_xlsx_upload_picks_the_test_case_sheet(tester):
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.active.append(["Connect Portal - Summary"])
    ws = wb.create_sheet("Test Cases")
    ws.append(HEAD)
    for r in ROWS[:3]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    p = mk_project(tester)
    out = preview(tester, p["id"], buf.getvalue(), name="cases.xlsx").json()
    assert out["sheet"] == "Test Cases" and out["valid"] == 3 and out["sheets"] == ["Summary", "Test Cases"]


def test_record_results_close_and_export(tester):
    p = mk_project(tester)
    commit(tester, p["id"], csv_bytes(), {"mapping": im.suggest_mapping(HEAD),
                                           "run_plan": {"mode": "single", "name": "All", "apply_results": False},
                                           "confirm": True})
    tr = tester.get(f"/api/projects/{p['id']}/test-runs").json()[0]
    assert tr["counts"]["untested"] == 3
    items = tester.get(f"/api/test-runs/{tr['id']}").json()["items"]
    r = tester.patch(f"/api/test-runs/{tr['id']}/items/{items[0]['id']}",
                     json={"status": "failed", "actual_result": "=cmd|' /C calc'!A0"})
    assert r.status_code == 200 and r.json()["executed_by"] == "tess"
    assert tester.post(f"/api/test-runs/{tr['id']}/items/bulk",
                       json={"item_ids": [i["id"] for i in items[1:]], "status": "passed"}).json()["updated"] == 2
    got = tester.get(f"/api/test-runs/{tr['id']}").json()
    assert got["progress"] == 100 and got["counts"]["failed"] == 1
    x = tester.get(f"/api/test-runs/{tr['id']}/export?format=xlsx")
    assert x.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    hdr = [c.value for c in ws[1]]
    assert ws.cell(2, hdr.index("Result") + 1).value == "Fail"
    assert ws.cell(2, hdr.index("Actual Result") + 1).value.startswith("'=")  # formula injection neutralised
    assert tester.patch(f"/api/test-runs/{tr['id']}", json={"status": "closed"}).json()["status"] == "closed"
    assert tester.patch(f"/api/test-runs/{tr['id']}/items/{items[0]['id']}", json={"status": "passed"}).status_code == 409


def test_create_runs_from_existing_cases(tester):
    p = mk_project(tester)
    commit(tester, p["id"], csv_bytes(), {"mapping": im.suggest_mapping(HEAD), "run_plan": {"mode": "none"},
                                           "confirm": True})
    plan = tester.post(f"/api/projects/{p['id']}/test-runs/plan",
                       json={"plan": {"mode": "group", "group_by": ["priority"]}}).json()
    assert {g["value"] for g in plan["groups"]} == {"High", "Medium"}
    made = tester.post(f"/api/projects/{p['id']}/test-runs",
                       json={"name": "Release 1.4", "plan": {"mode": "group", "group_by": ["priority"],
                                                            "groups": ["High"]}})
    assert made.status_code == 201, made.text
    assert [r["name"] for r in made.json()] == ["Release 1.4 – High"] and made.json()[0]["total"] == 2
    one = tester.post(f"/api/projects/{p['id']}/test-runs",
                      json={"name": "Dashboard only", "filter_field": "module", "filter_values": ["Dashboard"]}).json()
    assert one[0]["total"] == 1
    cases = tester.get(f"/api/projects/{p['id']}/cases").json()
    assert tester.post(f"/api/test-runs/{one[0]['id']}/items", json={"case_ids": [c["id"] for c in cases]}).json()["added"] == 2
    assert tester.delete(f"/api/test-runs/{one[0]['id']}").status_code == 200
    assert len(tester.get(f"/api/projects/{p['id']}/cases").json()) == 3  # cases survive


def test_viewer_cannot_import_or_record(viewer, tester):
    p = mk_project(tester)
    assert preview(viewer, p["id"], csv_bytes()).status_code == 403
    assert viewer.get("/api/imports/template?format=csv").status_code == 200


def test_automated_item_gets_worker_result(admin, sample_server):
    allow("127.0.0.1")
    p = admin.post("/api/projects", json={"name": "Sample", "base_url": sample_server}).json()
    auto = admin.post(f"/api/projects/{p['id']}/cases", json={"name": "Home", "steps": [
        {"action": "goto", "url": "/"}, {"action": "expect_visible", "selector": "body"}]}).json()
    admin.post(f"/api/projects/{p['id']}/cases", json={"name": "Manual only", "procedure": "1. Look."})
    tr = admin.post(f"/api/projects/{p['id']}/test-runs", json={"name": "Smoke"}).json()[0]
    out = admin.post(f"/api/test-runs/{tr['id']}/automate").json()
    assert out == {"queued": 1, "manual": 1, "problems": []}
    worker.run_once("pytest")
    items = {i["test_case_id"]: i for i in admin.get(f"/api/test-runs/{tr['id']}").json()["items"]}
    assert items[auto["id"]]["status"] == "passed" and items[auto["id"]]["automated_status"] == "passed"
    assert "TestAgent" in items[auto["id"]]["executed_by"]


def test_undo_import(tester):
    p = mk_project(tester)
    out = commit(tester, p["id"], csv_bytes(), {"mapping": im.suggest_mapping(HEAD),
                                                 "run_plan": {"mode": "group", "group_by": ["module"]},
                                                 "confirm": True}).json()
    u = tester.post(f"/api/projects/{p['id']}/imports/{out['batch']}/undo").json()
    assert u == {"archived_cases": 3, "deleted_test_runs": 2}
    assert tester.get(f"/api/projects/{p['id']}/cases").json() == []


# ---- multi-tenancy & project access (release 1.3 integration) ----------------------------
def test_test_runs_are_isolated_between_tenants(app):
    from .conftest import login, make_tenant

    make_tenant("acme")
    make_tenant("globex")
    a = login(app, "acme-admin", "tenant_admin", tenant="acme", platform=False)
    g = login(app, "globex-admin", "tenant_admin", tenant="globex", platform=False)
    pa = a.post("/api/t/acme/projects", json={"name": "A", "base_url": "https://a.example.com"}).json()
    pg = g.post("/api/t/globex/projects", json={"name": "G", "base_url": "https://g.example.com"}).json()
    opts = {"mapping": im.suggest_mapping(HEAD), "run_plan": {"mode": "single", "name": "A run"}, "confirm": True}
    r = a.post(f"/api/t/acme/projects/{pa['id']}/imports/commit", files={"file": ("c.csv", csv_bytes(), "text/csv")},
               data={"options": json.dumps(opts)})
    assert r.status_code == 201, r.text
    tr = r.json()["test_runs"][0]["id"]
    detail = a.get(f"/api/t/acme/test-runs/{tr}").json()
    iid = detail["items"][0]["id"]
    cid = detail["items"][0]["test_case_id"]
    # globex cannot see, change, export or extend acme's test run, nor import into acme's project
    assert g.get(f"/api/t/globex/test-runs/{tr}").status_code == 404
    assert g.patch(f"/api/t/globex/test-runs/{tr}/items/{iid}", json={"status": "failed"}).status_code == 404
    assert g.get(f"/api/t/globex/test-runs/{tr}/export").status_code == 404
    assert g.delete(f"/api/t/globex/test-runs/{tr}").status_code == 404
    assert g.get(f"/api/t/globex/projects/{pa['id']}/test-runs").status_code == 404
    assert g.post(f"/api/t/globex/projects/{pa['id']}/imports/preview",
                  files={"file": ("c.csv", csv_bytes(), "text/csv")}).status_code == 404
    # globex cannot pull acme's cases into its own test run
    gr = g.post(f"/api/t/globex/projects/{pg['id']}/test-runs", json={"name": "G run", "case_ids": [cid]})
    assert gr.status_code == 422  # no matching cases in globex's project
    assert g.get(f"/api/t/globex/projects/{pg['id']}/test-runs").json() == []
    # acme's own session on the wrong tenant's URL is refused
    assert a.get(f"/api/t/globex/test-runs/{tr}").status_code == 401
    assert a.get(f"/api/t/acme/test-runs/{tr}").json()["items"][0]["status"] == "passed"


def test_project_assignment_roles_gate_test_management(app):
    from .conftest import login, make_tenant

    make_tenant("acme")
    owner = login(app, "own", "tenant_owner", tenant="acme", platform=False)
    p = owner.post("/api/t/acme/projects", json={"name": "A", "base_url": "https://a.example.com"}).json()
    opts = {"mapping": im.suggest_mapping(HEAD), "run_plan": {"mode": "single", "name": "R"}, "confirm": True}
    tr = owner.post(f"/api/t/acme/projects/{p['id']}/imports/commit", files={"file": ("c.csv", csv_bytes(), "text/csv")},
                    data={"options": json.dumps(opts)}).json()["test_runs"][0]["id"]
    iid = owner.get(f"/api/t/acme/test-runs/{tr}").json()["items"][0]["id"]
    # an auditor with access to every project may read and export, not import or record results
    aud = login(app, "aud", "auditor", tenant="acme", platform=False)
    assert aud.get(f"/api/t/acme/test-runs/{tr}").status_code == 200
    assert aud.get(f"/api/t/acme/test-runs/{tr}/export?format=csv").status_code == 200
    assert aud.patch(f"/api/t/acme/test-runs/{tr}/items/{iid}", json={"status": "failed"}).status_code == 403
    assert aud.post(f"/api/t/acme/projects/{p['id']}/imports/preview",
                    files={"file": ("c.csv", csv_bytes(), "text/csv")}).status_code == 403
    # a developer without any assignment does not see the project at all
    dev = login(app, "dev", "developer", tenant="acme", platform=False, assign_all=False)
    assert dev.get(f"/api/t/acme/test-runs/{tr}").status_code == 404
    assert dev.get(f"/api/t/acme/projects/{p['id']}/test-runs").status_code == 404
