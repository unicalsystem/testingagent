"""CSV / Excel import of test cases, column mapping, run grouping, template and export.

Flow: the browser uploads the file to /imports/preview (nothing is written), the user checks
or changes the column mapping and chooses how to organise test runs, then confirms and the
same file is sent to /imports/commit. Parsing is pure and stateless so both calls agree.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from typing import Any, Iterable

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLS = 200
MAX_CELL = 20000


class ImportError_(ValueError):
    pass


# --------------------------------------------------------------------------- field registry
@dataclass(frozen=True)
class Field:
    key: str
    label: str            # canonical column header used by the template
    target: str           # "case" column | "attr" (case.attributes) | "result" (test-run item)
    synonyms: tuple[str, ...] = ()
    help: str = ""
    groupable: bool = False
    required: bool = False
    example: tuple[str, str] = ("", "")


FIELDS: list[Field] = [
    Field("external_id", "Test Case ID", "case", ("id", "test id", "tc id", "case id", "test case id", "test case no",
          "tc no", "test case number", "key", "test key", "ref", "reference"),
          "Your own ID. Re-importing a file with the same IDs updates those cases instead of duplicating them.",
          example=("LOG-001", "LOG-002")),
    Field("name", "Test Case Title", "case", ("title", "test case title", "test case name", "test case", "name",
          "test name", "summary", "test title", "scenario", "test scenario"),
          "Required. A short title of the check.", required=True,
          example=("Login with valid credentials", "Login with invalid password")),
    Field("description", "Description", "case", ("description", "objective", "purpose", "details"),
          "Optional longer description or objective."),
    Field("requirement", "Requirement", "case", ("requirement", "requirements", "requirement id", "req", "req id",
          "user story", "story", "traceability", "srs ref"), "Requirement / story reference (traceability).",
          groupable=True, example=("SRS-AUTH-01", "SRS-AUTH-01")),
    Field("module", "Module", "case", ("module", "feature", "area", "component", "functional area", "suite",
          "test suite", "section"), "Top-level area. A good way to group test runs.", groupable=True,
          example=("Login & Access", "Login & Access")),
    Field("sub_module", "Sub-module", "case", ("sub module", "submodule", "sub-module", "sub feature", "sub area",
          "screen", "page", "sub section", "subsection"), "Second-level area.", groupable=True,
          example=("Login", "Login")),
    Field("scenario_type", "Scenario Type", "case", ("scenario type", "type", "test type", "case type", "category"),
          "e.g. Positive, Negative, Boundary, Security/Access.", groupable=True, example=("Positive", "Negative")),
    Field("priority", "Priority", "case", ("priority", "prio", "importance"), "e.g. High, Medium, Low (or P1–P4).",
          groupable=True, example=("High", "High")),
    Field("test_level", "Test Level", "attr", ("test level", "level"), "e.g. Unit, Integration, System, Acceptance.",
          groupable=True, example=("System", "System")),
    Field("method", "Method", "attr", ("method", "verification method", "technique", "test method"),
          "e.g. Test, Inspection, Analysis, Demonstration.", groupable=True, example=("Test", "Test")),
    Field("preconditions", "Pre-conditions", "case", ("pre conditions", "preconditions", "pre condition",
          "precondition", "prerequisites", "prerequisite", "pre requisites", "setup"),
          "State required before the test starts.", example=("Valid active user account.", "Logged out.")),
    Field("test_data", "Test Data", "case", ("test data", "data", "input data", "inputs"), "Data used by the test.",
          example=("", "")),
    Field("procedure", "Procedure / Steps", "case", ("procedure steps", "procedure", "steps", "test steps",
          "step", "actions", "test procedure", "steps to execute"),
          "Manual steps, one per line (1. … 2. …).",
          example=("1. Open the login page.\n2. Enter valid username and password.\n3. Click Sign in.",
                   "1. Open the login page.\n2. Enter a wrong password.\n3. Click Sign in.")),
    Field("expected_result", "Expected Result", "case", ("expected result", "expected results", "expected",
          "expected outcome", "expected behaviour", "expected behavior"), "What must happen for a pass.",
          example=("User lands on the Dashboard.", "A generic error is shown; user stays on the login page.")),
    Field("tags", "Tags", "case", ("tags", "labels", "keywords", "label"), "Comma-separated labels.",
          example=("smoke, auth", "auth")),
    Field("assigned_to", "Assigned To", "case", ("assigned to", "assignee", "owner", "tester", "assigned"),
          "Person responsible for executing it.", groupable=True, example=("", "")),
    Field("version", "Version", "attr", ("version", "build", "release", "app version"), "Build or release under test.",
          groupable=True, example=("", "")),
    Field("source_project", "Project", "attr", ("project", "project name", "application"),
          "Informational only: cases are always imported into the project you are in.", groupable=True),
    Field("result", "Result", "result", ("result", "status", "execution status", "outcome", "test result",
          "pass fail", "verdict"),
          "Optional existing result: Pass, Fail, Blocked, Retest, Skipped or Not Run. Applied to the new test run(s).",
          example=("Not Run", "Not Run")),
    Field("actual_result", "Actual Result", "result", ("actual result", "actual", "actual results",
          "actual outcome", "observed result", "comments", "remarks"), "Optional observed result."),
    Field("executed_by", "Executed By", "result", ("executed by", "tested by", "run by", "executor"), ""),
    Field("executed_on", "Executed On", "result", ("executed on", "execution date", "tested on", "date executed",
          "run date", "test date", "executed date"), "Date, e.g. 2026-10-26."),
]
FIELD = {f.key: f for f in FIELDS}
GROUPABLE = [f.key for f in FIELDS if f.groupable]
CASE_TEXT_LIMITS = {"external_id": 100, "name": 200, "module": 200, "sub_module": 200, "priority": 40,
                    "scenario_type": 80, "assigned_to": 120}


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s).lower()).strip()


_SYN: dict[str, str] = {}
for _f in FIELDS:
    for _s in (_f.label, _f.key.replace("_", " "), *_f.synonyms):
        _SYN.setdefault(_norm(_s), _f.key)


def guess_field(header: str) -> str | None:
    n = _norm(header)
    if not n:
        return None
    if n in _SYN:
        return _SYN[n]
    stripped = _norm(re.sub(r"\(.*?\)|\[.*?\]", " ", str(header)))  # "Test Case ID (Doc)" -> "test case id"
    if stripped in _SYN:
        return _SYN[stripped]
    return None


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """header -> field key | "attr:<header>" (keep as custom field) | "" (ignore)."""
    used: set[str] = set()
    out: dict[str, str] = {}
    for h in headers:
        k = guess_field(h)
        if k and k not in used:
            out[h] = k
            used.add(k)
        else:
            out[h] = f"attr:{h}" if h.strip() else ""
    return out


# --------------------------------------------------------------------------- parsing
@dataclass
class Table:
    sheets: list[str]
    sheet: str
    headers: list[str]
    rows: list[list[str]]          # data rows (already trimmed to header width)
    first_row_number: int          # spreadsheet row number of rows[0] (for error messages)
    truncated: bool = False


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").strip()
    return s[:MAX_CELL]


def _score_header(row: list[str]) -> int:
    return sum(1 for c in row if guess_field(c))


def _pick_header(rows: list[list[str]]) -> int:
    best, best_score = 0, -1
    for i, r in enumerate(rows[:15]):
        sc = _score_header(r)
        if sc > best_score:
            best, best_score = i, sc
    if best_score <= 0:  # nothing recognised: first non-empty row
        best = next((i for i, r in enumerate(rows) if any(c.strip() for c in r)), 0)
    return best


def _finish(sheets: list[str], sheet: str, raw: list[list[str]]) -> Table:
    if not raw:
        raise ImportError_("The file has no rows")
    h = _pick_header(raw)
    headers = [c.strip() for c in raw[h]][:MAX_COLS]
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise ImportError_("No header row found")
    # unique, non-empty header names
    seen: dict[str, int] = {}
    for i, name in enumerate(headers):
        name = name or f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        headers[i] = name
    rows, truncated = [], False
    for r in raw[h + 1:]:
        r = (r + [""] * len(headers))[:len(headers)]
        rows.append(r)
    # drop trailing empty rows
    while rows and not any(c for c in rows[-1]):
        rows.pop()
    if len(rows) > MAX_ROWS:
        rows, truncated = rows[:MAX_ROWS], True
    return Table(sheets, sheet, headers, rows, h + 2, truncated)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportError_("Could not read the file's text encoding (save it as UTF-8 CSV)")


def parse_csv(data: bytes) -> Table:
    text = _decode(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    raw: list[list[str]] = []
    try:
        for i, r in enumerate(reader):
            if i > MAX_ROWS + 20:
                break
            raw.append([_cell(c) for c in r[:MAX_COLS]])
    except csv.Error as e:
        raise ImportError_(f"The CSV file is malformed: {e}")
    return _finish(["CSV"], "CSV", raw)


def _xlsx_rows(ws) -> list[list[str]]:
    raw = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i > MAX_ROWS + 20:
            break
        raw.append([_cell(c) for c in r[:MAX_COLS]])
    return raw


def parse_xlsx(data: bytes, sheet: str = "") -> Table:
    import openpyxl  # uses defusedxml when installed (it is pinned in requirements)

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ImportError_(f"Could not open the Excel workbook ({type(e).__name__}). Save it as .xlsx and retry.")
    try:
        names = wb.sheetnames
        if sheet:
            if sheet not in names:
                raise ImportError_(f"Sheet '{sheet}' not found")
            return _finish(names, sheet, _xlsx_rows(wb[sheet]))
        # choose the sheet that looks most like a test-case table
        best: tuple[int, int, str, list[list[str]]] | None = None
        for n in names:
            raw = _xlsx_rows(wb[n])
            if not raw:
                continue
            h = _pick_header(raw)
            score = (_score_header(raw[h]), len(raw))
            if best is None or score > best[:2]:
                best = (*score, n, raw)
        if best is None:
            raise ImportError_("The workbook has no data")
        return _finish(names, best[2], best[3])
    finally:
        wb.close()


def parse_file(filename: str, data: bytes, sheet: str = "") -> Table:
    if len(data) > MAX_BYTES:
        raise ImportError_(f"File is larger than {MAX_BYTES // (1024 * 1024)} MB")
    if not data:
        raise ImportError_("The file is empty")
    low = (filename or "").lower()
    if low.endswith((".xlsx", ".xlsm")) or data[:4] == b"PK\x03\x04":
        return parse_xlsx(data, sheet)
    if low.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ImportError_("Old .xls workbooks are not supported. In Excel use File → Save As → .xlsx (or CSV).")
    return parse_csv(data)


# --------------------------------------------------------------------------- normalisation
_RESULT = {
    "pass": "passed", "passed": "passed", "ok": "passed", "p": "passed", "success": "passed", "successful": "passed",
    "fail": "failed", "failed": "failed", "f": "failed", "failure": "failed", "nok": "failed", "not ok": "failed",
    "blocked": "blocked", "block": "blocked", "b": "blocked",
    "retest": "retest", "re test": "retest", "to retest": "retest",
    "skip": "skipped", "skipped": "skipped", "na": "skipped", "n a": "skipped", "not applicable": "skipped",
    "not run": "untested", "notrun": "untested", "untested": "untested", "not tested": "untested", "": "untested",
    "no run": "untested", "pending": "untested", "not executed": "untested", "to do": "untested", "todo": "untested",
}
_PRIO = {"high": "High", "medium": "Medium", "med": "Medium", "low": "Low", "critical": "Critical",
         "urgent": "Critical", "blocker": "Critical", "normal": "Medium", "major": "High", "minor": "Low"}


def norm_result(v: str) -> str | None:
    return _RESULT.get(_norm(v))


def norm_priority(v: str) -> str:
    n = _norm(v)
    if n in _PRIO:
        return _PRIO[n]
    m = re.fullmatch(r"p ?([0-9])", n)
    return f"P{m.group(1)}" if m else v.strip()


def parse_date(v: str) -> datetime | None:
    v = (v or "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y",
                "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%m/%d/%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(v, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def split_tags(v: str) -> list[str]:
    return [t.strip()[:40] for t in re.split(r"[,;|]", v or "") if t.strip()][:20]


@dataclass
class Record:
    row: int
    case: dict[str, Any]           # TestCase column values
    attributes: dict[str, str]
    result: dict[str, str]
    warnings: list[str] = field(default_factory=list)


@dataclass
class Mapped:
    records: list[Record]
    errors: list[dict]             # rows that cannot be imported
    warnings: list[dict]
    in_file_duplicates: int


def validate_mapping(headers: list[str], mapping: dict[str, str]) -> dict[str, str]:
    clean: dict[str, str] = {}
    used: dict[str, str] = {}
    for h in headers:
        v = (mapping.get(h) or "").strip()
        if v.startswith("attr:"):
            clean[h] = "attr:" + (v[5:].strip() or h)[:80]
            continue
        if v and v not in FIELD:
            raise ImportError_(f"Unknown field '{v}' for column '{h}'")
        if v and v in used:
            raise ImportError_(f"Columns '{used[v]}' and '{h}' are both mapped to {FIELD[v].label}")
        if v:
            used[v] = h
        clean[h] = v
    if "name" not in used and "external_id" not in used:
        raise ImportError_("Map a column to Test Case Title (or at least to Test Case ID)")
    return clean


def apply_mapping(table: Table, mapping: dict[str, str]) -> Mapped:
    mapping = validate_mapping(table.headers, mapping)
    idx = {h: i for i, h in enumerate(table.headers)}
    records: list[Record] = []
    errors: list[dict] = []
    warnings: list[dict] = []
    seen_ids: dict[str, int] = {}
    dups = 0
    for n, r in enumerate(table.rows):
        rownum = table.first_row_number + n
        if not any(c.strip() for c in r):
            continue
        case: dict[str, Any] = {}
        attrs: dict[str, str] = {}
        res: dict[str, str] = {}
        for h, key in mapping.items():
            if not key:
                continue
            val = r[idx[h]].strip()
            if key.startswith("attr:"):
                if val:
                    attrs[key[5:]] = val[:2000]
                continue
            f = FIELD[key]
            if f.target == "attr":
                if val:
                    attrs[f.label] = val[:2000]
            elif f.target == "result":
                res[key] = val
            elif key == "tags":
                case["tags"] = split_tags(val)
            else:
                case[key] = val
        rec_warn: list[str] = []
        for k, lim in CASE_TEXT_LIMITS.items():
            if len(case.get(k, "")) > lim:
                case[k] = case[k][:lim]
                rec_warn.append(f"{FIELD[k].label} shortened to {lim} characters")
        if case.get("priority"):
            case["priority"] = norm_priority(case["priority"])[:40]
        if not case.get("name"):
            if case.get("external_id"):
                case["name"] = case["external_id"]
                rec_warn.append("no title: the ID is used as the title")
            else:
                errors.append({"row": rownum, "error": "No title and no ID"})
                continue
        if "result" in res:
            st = norm_result(res["result"])
            if st is None:
                rec_warn.append(f"unknown result '{res['result'][:30]}' treated as Not Run")
                st = "untested"
            res["status"] = st
        ext = case.get("external_id", "")
        if ext:
            key_ = ext.lower()
            if key_ in seen_ids:
                dups += 1
                errors.append({"row": rownum, "error": f"Duplicate ID {ext} (first seen on row {seen_ids[key_]})"})
                continue
            seen_ids[key_] = rownum
        if rec_warn:
            warnings.append({"row": rownum, "warning": "; ".join(rec_warn)})
        records.append(Record(rownum, case, attrs, res, rec_warn))
    return Mapped(records, errors, warnings, dups)


# --------------------------------------------------------------------------- grouping
def group_value(source: dict[str, Any], attrs: dict[str, Any], key: str) -> str:
    """Value of a groupable field for a case dict (or a Record.case) — "" if empty."""
    if key.startswith("attr:"):
        return str(attrs.get(key[5:], "") or "").strip()
    f = FIELD.get(key)
    if not f:
        return ""
    if f.target == "attr":
        return str(attrs.get(f.label, "") or "").strip()
    return str(source.get(key, "") or "").strip()


def group_label(keys: list[str]) -> str:
    return " / ".join(FIELD[k].label if k in FIELD else k[5:] for k in keys)


def group_key(source: dict, attrs: dict, keys: list[str]) -> str:
    return " / ".join(group_value(source, attrs, k) or "(none)" for k in keys)


def groups_for(items: Iterable[tuple[dict, dict]], keys: list[str]) -> list[dict]:
    counts: dict[str, int] = {}
    for src, attrs in items:
        g = group_key(src, attrs, keys)
        counts[g] = counts.get(g, 0) + 1
    return [{"value": k, "count": v} for k, v in counts.items()]  # insertion order = file order


def group_options(records: list[Record], custom_attrs: list[str]) -> list[dict]:
    out = []
    keys = GROUPABLE + [f"attr:{a}" for a in custom_attrs]
    for k in keys:
        vals = {group_value(r.case, r.attributes, k) for r in records}
        vals.discard("")
        if not vals:
            continue
        out.append({"field": k, "label": FIELD[k].label if k in FIELD else k[5:], "distinct": len(vals)})
    return out


def recommend_grouping(options: list[dict], total: int) -> list[str]:
    by = {o["field"]: o["distinct"] for o in options}
    for k in ("module", "sub_module", "requirement", "scenario_type", "priority"):
        if 2 <= by.get(k, 0) <= 40:
            return [k]
    return []


# --------------------------------------------------------------------------- template & export
TEMPLATE_FIELDS = ["external_id", "name", "description", "requirement", "module", "sub_module", "scenario_type",
                   "priority", "test_level", "method", "preconditions", "test_data", "procedure", "expected_result",
                   "tags", "assigned_to", "version", "result", "actual_result", "executed_by", "executed_on"]


def template_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([FIELD[k].label for k in TEMPLATE_FIELDS])
    for i in range(2):
        w.writerow([FIELD[k].example[i] for k in TEMPLATE_FIELDS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _xlsx_header(ws, labels: list[str], widths: dict[int, int] | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.append(labels)
    fill = PatternFill("solid", fgColor="14305C")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    from openpyxl.utils import get_column_letter

    for i, lab in enumerate(labels, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(14, min(len(lab) + 4, 40)))


def template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    labels = [FIELD[k].label for k in TEMPLATE_FIELDS]
    wide = {TEMPLATE_FIELDS.index(k) + 1: 45 for k in ("name", "preconditions", "procedure", "expected_result")}
    _xlsx_header(ws, labels, wide)
    for i in range(2):
        ws.append([FIELD[k].example[i] for k in TEMPLATE_FIELDS])
    for row in ws.iter_rows(min_row=2, max_row=3):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    from openpyxl.utils import get_column_letter

    def dv(key: str, values: list[str]) -> None:
        col = get_column_letter(TEMPLATE_FIELDS.index(key) + 1)
        v = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        v.error, v.showErrorMessage = "Pick a value from the list", False  # warn only: other values are allowed
        ws.add_data_validation(v)
        v.add(f"{col}2:{col}{MAX_ROWS + 1}")

    dv("priority", ["Critical", "High", "Medium", "Low"])
    dv("scenario_type", ["Positive", "Negative", "Boundary", "Security/Access", "Usability", "Performance"])
    dv("result", ["Not Run", "Pass", "Fail", "Blocked", "Retest", "Skipped"])

    info = wb.create_sheet("Instructions")
    _xlsx_header(info, ["Column", "Required", "What to enter"], {1: 22, 2: 10, 3: 100})
    info.append(["", "", "Fill one row per test case on the 'Test Cases' sheet. Delete the two example rows. "
                 "Column names can differ: TestAgent lets you match each column after upload."])
    for k in TEMPLATE_FIELDS:
        f = FIELD[k]
        info.append([f.label, "Yes" if f.required else "", f.help])
    for row in info.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


STATUS_LABEL = {"untested": "Not Run", "passed": "Pass", "failed": "Fail", "blocked": "Blocked",
                "retest": "Retest", "skipped": "Skipped"}
EXPORT_HEADERS = [FIELD[k].label for k in TEMPLATE_FIELDS] + ["Notes", "Automated Run"]


def export_rows(items: list[Any]) -> list[list[str]]:
    rows = []
    for it in items:
        c = it.case
        a = c.attributes or {}
        vals = {
            "external_id": c.external_id, "name": c.name, "description": c.description, "requirement": c.requirement,
            "module": c.module, "sub_module": c.sub_module, "scenario_type": c.scenario_type, "priority": c.priority,
            "test_level": a.get("Test Level", ""), "method": a.get("Method", ""), "preconditions": c.preconditions,
            "test_data": c.test_data, "procedure": c.procedure, "expected_result": c.expected_result,
            "tags": ", ".join(c.tags or []), "assigned_to": it.assigned_to or c.assigned_to,
            "version": a.get("Version", ""), "result": STATUS_LABEL.get(it.status, it.status),
            "actual_result": it.actual_result, "executed_by": it.executed_by_name,
            "executed_on": it.executed_at.date().isoformat() if it.executed_at else "",
        }
        rows.append([vals[k] or "" for k in TEMPLATE_FIELDS] + [it.notes or "",
                                                                  f"#{it.automated_run_id}" if it.automated_run_id else ""])
    return rows


def _safe_cell(v: str) -> str:
    # Spreadsheet formula injection: neutralise cells that a spreadsheet would evaluate.
    return "'" + v if v and v[0] in "=+-@\t\r" else v


def export_csv(items: list[Any]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXPORT_HEADERS)
    for r in export_rows(items):
        w.writerow([_safe_cell(v) for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_xlsx(items: list[Any], title: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    wide = {TEMPLATE_FIELDS.index(k) + 1: 45 for k in ("name", "procedure", "expected_result", "actual_result")}
    _xlsx_header(ws, EXPORT_HEADERS, wide)
    for r in export_rows(items):
        ws.append([_safe_cell(v) for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
